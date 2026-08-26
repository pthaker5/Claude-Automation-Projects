# =============================================================================
# Akzo Aging Pipeline
# -----------------------------------------------------------------------------
# Sources:
#   SOA  : GP_Lakehouse (Fabric)  - GP.vw_Open_Invoices, browser (MFA) login
#   TMW  : az-bwprod / CLXDW       - 712 billing extract, SQL login
#   IB   : H: drive Excel          - IB US/CAD detail (BU + Statement #)
#   BU   : H: drive Excel          - Origin location -> Business Unit
#
# Outputs:
#   1. Combined_Aging_Sources_Python_Output.xlsx  - full working file (internal)
#   2. Akzo_Aging_Report_for_Maria.xlsx           - dashboard-accurate summary
#      (emailed to Maria; matches the Tableau aging dashboard's Billed/Unbilled)
#
# Run flags:
#   (none)       build both files, email the Maria report immediately
#   --display    build both files, create an email DRAFT for manual review/send
#   --no-email   build both files, send nothing
#
# Email: sent through classic Outlook via COM automation - no IT setup
# needed. New Outlook (olk.exe) has no COM interface, so the script attaches
# to a running classic Outlook (starting one itself if needed), then forces
# a send/receive and waits for the Outbox to drain so the message actually
# transmits even while you work in new Outlook.
# =============================================================================

import os
import sys
import time
import struct
import getpass
import warnings
import subprocess
from pathlib import Path

import pandas as pd
import pyodbc
from azure.identity import InteractiveBrowserCredential

warnings.filterwarnings("ignore", category=UserWarning)

# =============================================================================
# CONFIG
# =============================================================================

# H: drive inputs
BU_LOOKUP_PATH = Path(r"H:\Integrated Logistics Design\Akzo Performance Coatings\Poojan Transition\Lookups for BU and Interplant\Akzo Origins – BU (09.04.2025).xlsx")
IB_CAD_2025 = Path(r"H:\TMC Customers\Akzo 4PL Coatings\Freight Payment\Akzo Coatings Billing File\IB US_CAD Combined YTD 2025\IB US_CAD YTD 2025.xlsx")
IB_CAD_2026 = Path(r"H:\TMC Customers\Akzo 4PL Coatings\Freight Payment\Akzo Coatings Billing File\IB US_CAD Combined YTD 2026\IB US_CAD YTD 2026.xlsx")

# Outputs
BASE_DIR = Path(r"H:\Integrated Logistics Design\Akzo Performance Coatings\Akzo Aging Reports")
AGING_DIR = BASE_DIR / "Aging 2025"
OUTPUT_PATH = AGING_DIR / "Combined_Aging_Sources_Python_Output.xlsx"
MARIA_REPORT_PATH = AGING_DIR / "Akzo_Aging_Report_for_Maria.xlsx"

# Fabric (GP_Lakehouse) - SOA source
FABRIC_SERVER = "gezt5um5sjaerdz52suwrkrx3q-m6nlihc5lfou3jj5p6tf7rptim.datawarehouse.fabric.microsoft.com"
FABRIC_DATABASE = "GP_Lakehouse"

# TMW 712 - az-bwprod (SQL login)
SQL_TABLE = "dbo.TMSCL712_3_FI_Billing_Extract_Akzo"
SQL_USER = os.environ.get("SQL_USER", "")
SQL_PASS = os.environ.get("SQL_PASS", "")
if not SQL_USER:
    SQL_USER = input("SQL Username (az-bwprod): ")
    SQL_PASS = getpass.getpass("SQL Password: ")

SQL_CONN_STRING = (
    "DRIVER={ODBC Driver 18 for SQL Server};"
    "SERVER=az-bwprod.chemlogix.com;"
    "DATABASE=CLXDW;"
    f"UID={SQL_USER};PWD={SQL_PASS};"
    "TrustServerCertificate=yes;"
)

# Email
EMAIL_TO = "mbates@quantixscs.com"
EMAIL_ENABLED = "--no-email" not in sys.argv
EMAIL_DRAFT_ONLY = "--display" in sys.argv

# BU normalization
BU_TO_FINAL_BU = {
    "POWDER": "Powder", "Powder": "Powder",
    "Wood": "Wood", "WOOD": "Wood", "WOOD ": "Wood",
    "MPY": "M & PC", "MPC": "M & PC", "MPC ": "M & PC",
    "METAL": "Metal", "Metal": "Metal", "METL": "Metal",
    "VR/ Specialty": "VR", "VR": "VR", "SC": "VR",
    "null": "null", "0": "null",
}


def clean_key(val):
    s = str(val).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s


def isblank(series):
    return series.isna() | series.astype(str).str.strip().isin(["", "nan", "None"])


# =============================================================================
# LOAD SOURCES
# =============================================================================

def load_soa(credential):
    """SOA (open invoices) from GP_Lakehouse via browser-auth token."""
    print("\nConnecting to GP_Lakehouse (browser login will open)...")
    token = credential.get_token("https://database.windows.net/.default")
    tb = token.token.encode("UTF-16-LE")
    token_struct = struct.pack(f"<I{len(tb)}s", len(tb), tb)
    conn = pyodbc.connect(
        "DRIVER={ODBC Driver 18 for SQL Server};"
        f"SERVER=tcp:{FABRIC_SERVER},1433;DATABASE={FABRIC_DATABASE};"
        "Encrypt=yes;TrustServerCertificate=no;",
        attrs_before={1256: token_struct},
    )
    q = """
    SELECT [Customer Number], [Customer Name], [Document Number], [Document Date],
           [Due Date], [Days Past Due], [Sales Amount], [Current Trx Amount],
           [Customer PO Number], [Document Type], [Document Description]
    FROM [GP].[vw_Open_Invoices]
    WHERE [Customer Number] IN ('AKZOLONY', 'AKZOLCAD')
    """
    print("Pulling SOA from GP_Lakehouse...")
    soa = pd.read_sql(q, conn)
    conn.close()
    print(f"  SOA rows: {len(soa):,}")

    soa["Document Number"] = soa["Document Number"].astype(str).str.strip()
    soa["Updated Document Number"] = soa["Document Number"].apply(
        lambda s: s[3:] if s.upper().startswith("QMS") else s)
    soa["Document Date"] = pd.to_datetime(soa["Document Date"], errors="coerce")
    soa["Due Date"] = pd.to_datetime(soa["Due Date"], errors="coerce")
    soa["Days Past Due"] = pd.to_numeric(soa["Days Past Due"], errors="coerce").fillna(0).astype(int)
    soa["Document Description"] = soa["Document Description"].astype(str).str.strip()
    return soa


def load_tmw(soa):
    """712 billing extract, filtered to the SOA SIDs, batched to avoid IN-list limits."""
    print("\nQuerying TMW (az-bwprod), filtered to SOA SIDs...")
    t0 = time.time()
    conn = pyodbc.connect(SQL_CONN_STRING, timeout=60)
    sids = [s for s in soa["Document Description"].dropna().astype(str).unique().tolist()
            if s and s != "nan"]
    print(f"  {len(sids)} unique SIDs")

    frames = []
    for i in range(0, len(sids), 500):
        batch = sids[i:i + 500]
        ph = ",".join("?" for _ in batch)
        q = f"""
            SELECT [SID], [Order Number], [Carrier Pro No], [Carrier SCAC], [Carrier Name],
                   [Movement Type], [Pick Up Date], [Delivery Date],
                   [Origin Loc Code], [Origin Name], [Origin City], [Or State],
                   [Dest Loc Code], [Destination Name], [Destination City], [Dest State],
                   [Bill To No], [Ship't Actual Cost], [Fuel Charges], [Currency Indicator],
                   [Shipment Status], [Service Type], [Equipment Type], [Transport Mode],
                   [User Name], [Created Date], [Billing Account], [Carrier Invoice Number],
                   [Carrier PRO Nbr], [Customer Reference Number], [Origin Country], [Dest Country]
            FROM {SQL_TABLE}
            WHERE [SID] IN ({ph})
        """
        frames.append(pd.read_sql(q, conn, params=batch))
        print(f"  Batch {i // 500 + 1}: {len(frames[-1])} rows")
    conn.close()

    tmw = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()
    print(f"  TMW total: {len(tmw)} rows ({time.time() - t0:.1f}s)")
    return tmw.drop_duplicates(subset="SID", keep="first").set_index("SID")


def load_ib():
    """IB US/CAD detail -> BU map and Statement # map, keyed by cleaned CLX invoice #."""
    print("\nReading IB US/CAD detail files...")
    ib = pd.concat([
        pd.read_excel(IB_CAD_2025, sheet_name="US_CAD IB Detail"),
        pd.read_excel(IB_CAD_2026, sheet_name="US_CAD IB Detail"),
    ], ignore_index=True)
    ib["_key"] = ib["CLX Invoice Number"].apply(clean_key)
    ded = ib.drop_duplicates(subset=["_key"], keep="first")
    bu_map = dict(zip(ded["_key"], ded["Business Unit"]))
    stmt_map = dict(zip(ded["_key"], ded["Statement #"]))
    print(f"  IB lookup keys: {len(bu_map)}")
    return bu_map, stmt_map


def load_bu_lookup():
    print("\nReading Business Unit lookup...")
    df = pd.read_excel(BU_LOOKUP_PATH)
    m = dict(zip(df.iloc[:, 0].astype(str).str.strip(),
                 df.iloc[:, 1].astype(str).str.strip()))
    print(f"  {len(m)} location-to-BU mappings")
    return m


# =============================================================================
# BUILD AGING DATASET
# =============================================================================

def build_aging(soa, tmw, ib_bu_map, ib_stmt_map, bu_lookup):
    print("\nBuilding aging dataset...")

    def tmw_col(col):
        return soa["Document Description"].map(tmw[col]).fillna("") if col in tmw.columns \
            else pd.Series("", index=soa.index)

    r = pd.DataFrame()
    r["BILL NUMBER"] = soa["Updated Document Number"].values
    r["INVOICE AMOUNT"] = soa["Current Trx Amount"].values
    r["BILL DATE"] = soa["Document Date"].values
    r["SID"] = soa["Document Description"].values
    r["BALANCE"] = soa["Current Trx Amount"].values

    sid = soa["Document Description"]
    r["CARRIER INVOICE #"] = tmw_col("Carrier Invoice Number").values
    r["CARRIER INVOICE DATE"] = ""
    r["ORIG ID"] = tmw_col("Origin Loc Code").values
    r["SCAC"] = tmw_col("Carrier SCAC").values
    r["Destination ID"] = tmw_col("Dest Loc Code").values
    r["Origin City"] = tmw_col("Origin City").values
    r["Destination City"] = tmw_col("Destination City").values
    r["Pick Up Date"] = tmw_col("Pick Up Date").values
    r["Carrier Name"] = tmw_col("Carrier Name").values

    # Movement Type: Inbound if SID starts AK, else TMW value, else Outbound
    tmw_mv = sid.map(tmw["Movement Type"]) if "Movement Type" in tmw.columns \
        else pd.Series("", index=sid.index)
    mv = sid.apply(lambda s: "Inbound" if str(s).startswith("AK") else "")
    not_ib = mv != "Inbound"
    mv[not_ib] = tmw_mv[not_ib].fillna("Outbound").values
    r["Movement Type"] = mv.values

    r["DAYS PAST"] = soa["Days Past Due"].values
    for col in ["STERLING VOUCHER #", "SAP ID", "BUSINESS UNIT", "AGE",
                "0 - 30", "31 - 60", "61 - 90", "91 - 120", "121+",
                "Pmt - Sep", "Pmt - Oct", "Column23"]:
        r[col] = ""

    # --- BU cascade + Statement # ---
    bill_clean = r["BILL NUMBER"].apply(clean_key)
    r["712 BU"] = r["ORIG ID"].astype(str).str.strip().map(bu_lookup)
    no712 = r["712 BU"].isna()
    r.loc[no712, "712 BU"] = r.loc[no712, "Destination ID"].astype(str).str.strip().map(bu_lookup)
    r["IB BU"] = bill_clean.map(ib_bu_map)
    r["Statement Number"] = bill_clean.map(ib_stmt_map).fillna("")
    r["STMT NUM"] = r["Statement Number"]
    r["MPC BU"] = pd.NA
    r["Voucher File BU"] = pd.NA
    r["BU"] = r["IB BU"].fillna(r["712 BU"]).fillna("null")
    r["Final BU"] = r["BU"].astype(str).str.strip().map(BU_TO_FINAL_BU).fillna("null")
    print(f"  712 BU: {r['712 BU'].notna().sum()} | IB BU: {r['IB BU'].notna().sum()} | "
          f"Stmt#: {(r['Statement Number'] != '').sum()} of {len(r)}")

    # --- Derived fields ---
    r["Bill Prefix"] = r["BILL NUMBER"].astype(str).apply(
        lambda x: "CA" if x[:2].upper() == "CA" else "Other")
    r["Movement Check"] = r["Movement Type"].apply(
        lambda x: "Outbound/Interplant" if x in ("Outbound", "Interplant") else "Inbound")
    r["SCAC2"] = r["SCAC"].apply(lambda x: "NO SCAC" if str(x).strip() == "" else x)
    r["Currency"] = r["BILL NUMBER"].astype(str).apply(
        lambda x: "CAD" if x[:2].upper() == "CA" else "USD")
    r["Due Date"] = soa["Due Date"].values
    r["Overdue"] = r["DAYS PAST"].apply(lambda x: "Y" if int(x) > 0 else "N")
    r["Not Yet Billed"] = ""

    cutoff = pd.Timestamp("2024-09-05")

    def calc_remove(x):
        try:
            if pd.notna(x) and x != "" and pd.Timestamp(x) < cutoff:
                return "Remove"
        except Exception:
            pass
        return ""
    r["Remove"] = r["Pick Up Date"].apply(calc_remove)
    r["File Number"] = ""

    def manual_voucher(row):
        s = str(row["SID"]).strip()
        if s in ("", "NONE") or s.startswith("AK"):
            return "Manual"
        fn = str(row.get("File Number", "")).strip()
        sn = str(row.get("Statement Number", "")).strip()
        if fn in ("", "nan") and sn in ("", "nan"):
            return "Voucher"
        if fn in ("", "0") and sn in ("", "0"):
            return "Voucher"
        return "Manual"
    r["Manual/Voucher"] = r.apply(manual_voucher, axis=1)
    r["Sum Row"] = r["BILL NUMBER"].apply(
        lambda x: "sum-remove" if str(x).strip() == "" else "")

    def billing_location(row):
        fbu = str(row["Final BU"]).strip()
        is_ca = str(row["BILL NUMBER"]).strip()[:2].upper() == "CA"
        mt = str(row["Movement Type"]).strip()
        oc = str(row["Origin City"]).strip()
        dc = str(row["Destination City"]).strip()
        if fbu == "VR":
            return "VR & SC"
        if fbu == "M & PC":
            return "M & PC CAN" if is_ca else "M & PC USA"
        if fbu == "Powder":
            return "Powder - CAN" if is_ca else "Powder - USA"
        if fbu == "Metal":
            return oc if mt in ("Outbound", "Interplant") else dc
        if fbu == "Wood":
            region = "CAN" if is_ca else "USA"
            leg = {"Inbound": "IB", "Outbound": "OB", "Interplant": "IP"}.get(mt)
            if leg:
                return f"Wood - {leg} - {region}"
        return "null"
    r["Billing Location"] = r.apply(billing_location, axis=1)

    final_columns = [
        "BILL NUMBER", "INVOICE AMOUNT", "BILL DATE", "CARRIER INVOICE #",
        "STMT NUM", "CARRIER INVOICE DATE", "ORIG ID", "DAYS PAST",
        "STERLING VOUCHER #", "SCAC", "SID", "SAP ID", "BUSINESS UNIT",
        "AGE", "BALANCE", "0 - 30", "31 - 60", "61 - 90", "91 - 120", "121+",
        "Pmt - Sep", "Pmt - Oct", "Column23",
        "Destination ID", "BU", "712 BU", "MPC BU", "IB BU",
        "Voucher File BU", "Final BU", "Movement Type", "Billing Location",
        "Origin City", "Destination City", "Bill Prefix", "Movement Check",
        "SCAC2", "File Number", "Statement Number", "Manual/Voucher",
        "Sum Row", "Due Date", "Overdue", "Not Yet Billed", "Pick Up Date",
        "Carrier Name", "Remove", "Currency",
    ]
    for col in final_columns:
        if col not in r.columns:
            r[col] = ""
    out = r[final_columns].copy()
    print(f"  Aging dataset: {len(out)} rows x {len(out.columns)} cols")
    return out


# =============================================================================
# MARIA REPORT (dashboard-accurate + readable)
# =============================================================================

def compute_status(d):
    """Billed/UnBilled per the Tableau workbook calc."""
    def _billed(row):
        bg = str(row["Final BU"]).strip()
        stmt = str(row.get("Statement Number", "")).strip()
        mvv = str(row["Manual/Voucher"]).strip()
        sid = str(row["SID"]).strip()
        ns = stmt in ("", "nan", "None", "0")
        if bg == "M & PC" and ns and mvv == "Voucher":
            return "UnBilled"
        if bg == "Wood" and ns and mvv == "Voucher" and sid[:2].upper() != "CA":
            return "UnBilled"
        return "Billed"
    return d.apply(_billed, axis=1)


def dashboard_base(d):
    """Apply the dashboard's filters: Remove not set, SCAC present, no sum rows."""
    keep = (
        (d["Remove"].astype(str).str.strip() != "Remove")
        & (~isblank(d["SCAC"]))
        & (d["Sum Row"].astype(str).str.strip() != "sum-remove")
    )
    return d[keep].copy()


def _detail_frame(base):
    """Row-level detail Maria can VLOOKUP on SID or FB#/Bill Number."""
    detail_cols = [
        "SID", "BILL NUMBER", "STMT NUM", "Status", "BALANCE",
        "Final BU", "Billing Location", "Movement Type", "BILL DATE",
        "Due Date", "Overdue", "SCAC", "Carrier Name", "Manual/Voucher",
        "Statement Number", "Currency",
    ]
    detail_cols = [c for c in detail_cols if c in base.columns]
    det = (base[detail_cols]
           .rename(columns={"BILL NUMBER": "FB# (Bill Number)", "Status": "Billed/Unbilled"}))
    return det.sort_values(["Billed/Unbilled", "Final BU", "SID"]).reset_index(drop=True)


def build_maria_report(src, out_path):
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    NAVY = PatternFill("solid", fgColor="1F4E78")
    HILITE = PatternFill("solid", fgColor="D9E1F2")
    MONEY = "#,##0.00"

    d = src.copy()
    d["_bal"] = pd.to_numeric(d["BALANCE"], errors="coerce").fillna(0)
    d["Status"] = compute_status(d)
    base = dashboard_base(d)

    billed = base[base["Status"] == "Billed"]
    unbilled = base[base["Status"] == "UnBilled"]

    def loc_summary(sub):
        return (sub.groupby("Billing Location")
                   .agg(Balance=("_bal", "sum"), Invoices=("BILL NUMBER", "nunique"))
                   .sort_values("Balance", ascending=False).reset_index())

    billed_sum = loc_summary(billed)
    unbilled_sum = loc_summary(unbilled)
    detail = _detail_frame(base)

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        billed_sum.to_excel(writer, sheet_name="Billed - by Location", index=False, startrow=2)
        unbilled_sum.to_excel(writer, sheet_name="Unbilled - by Location", index=False, startrow=2)
        detail.to_excel(writer, sheet_name="Detail (Dashboard Rows)", index=False)
        wb = writer.book

        # ---- Overview ----
        ov = wb.create_sheet("Overview", 0)
        ov["A1"] = "Akzo Aging Summary"
        ov["A1"].font = Font(size=16, bold=True, color="1F4E78")
        ov["A2"] = f"As of {pd.Timestamp.now():%B %d, %Y}"
        ov["A2"].font = Font(italic=True, color="595959")

        rows = [
            ("", "Balance", "Open Invoices"),
            ("Billed", billed["_bal"].sum(), billed["BILL NUMBER"].nunique()),
            ("Unbilled", unbilled["_bal"].sum(), unbilled["BILL NUMBER"].nunique()),
            ("Total Outstanding", base["_bal"].sum(), base["BILL NUMBER"].nunique()),
        ]
        for i, (a, b, c) in enumerate(rows):
            rr = 4 + i
            ov.cell(rr, 1, a)
            ov.cell(rr, 2, b)
            ov.cell(rr, 3, c)
            if i == 0:
                for col in (1, 2, 3):
                    cell = ov.cell(rr, col)
                    cell.font = Font(size=12, bold=True, color="FFFFFF")
                    cell.fill = NAVY
                    cell.alignment = Alignment(horizontal="center")
            else:
                ov.cell(rr, 1).font = Font(bold=True)
                ov.cell(rr, 2).number_format = MONEY
                if a == "Total Outstanding":
                    for col in (1, 2, 3):
                        ov.cell(rr, col).font = Font(bold=True, size=12)
                        ov.cell(rr, col).fill = HILITE
        ov.column_dimensions["A"].width = 22
        ov.column_dimensions["B"].width = 18
        ov.column_dimensions["C"].width = 16
        ov["A10"] = "Notes"
        ov["A10"].font = Font(bold=True)
        ov["A11"] = "Figures match the Akzo Aging Tableau dashboard (billed vs unbilled)."
        ov["A12"] = "Balance = open dollars outstanding.  Open Invoices = distinct bill numbers."
        ov["A13"] = "Detail (Dashboard Rows) tab = every record on the dashboard; VLOOKUP on SID or FB#."
        for c in ("A11", "A12", "A13"):
            ov[c].font = Font(color="595959")

        # ---- Location sheets ----
        for sh, gdf, heading in [
            ("Billed - by Location", billed_sum, "Billed - Balance by Billing Location"),
            ("Unbilled - by Location", unbilled_sum, "Unbilled - Balance by Billing Location"),
        ]:
            ws = wb[sh]
            ws["A1"] = heading
            ws["A1"].font = Font(size=13, bold=True, color="1F4E78")
            for col in range(1, 4):
                cell = ws.cell(3, col)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = NAVY
                cell.alignment = Alignment(horizontal="center")
            n = len(gdf)
            for rr in range(4, 4 + n):
                ws.cell(rr, 2).number_format = MONEY
            tr = 4 + n
            ws.cell(tr, 1, "Grand Total").font = Font(bold=True)
            gt = ws.cell(tr, 2, float(gdf["Balance"].sum()))
            gt.number_format = MONEY
            gt.font = Font(bold=True)
            ws.cell(tr, 3, int(gdf["Invoices"].sum())).font = Font(bold=True)
            for col in (1, 2, 3):
                ws.cell(tr, col).fill = HILITE
            ws.column_dimensions["A"].width = 26
            ws.column_dimensions["B"].width = 16
            ws.column_dimensions["C"].width = 12

        # ---- Detail sheet (VLOOKUP-ready) ----
        wd = wb["Detail (Dashboard Rows)"]
        wd.freeze_panes = "A2"
        wd.auto_filter.ref = wd.dimensions
        for ci, _ in enumerate(detail.columns, 1):
            cell = wd.cell(1, ci)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = NAVY
            cell.alignment = Alignment(horizontal="center")
        if "BALANCE" in detail.columns:
            bal_i = list(detail.columns).index("BALANCE") + 1
            for rr in range(2, len(detail) + 2):
                wd.cell(rr, bal_i).number_format = MONEY
        det_widths = {
            "SID": 16, "FB# (Bill Number)": 18, "STMT NUM": 14,
            "Billed/Unbilled": 15, "BALANCE": 14, "Final BU": 12,
            "Billing Location": 20, "Movement Type": 14, "BILL DATE": 12,
            "Due Date": 12, "Overdue": 9, "SCAC": 8, "Carrier Name": 24,
            "Manual/Voucher": 15, "Statement Number": 16, "Currency": 9,
        }
        for ci, col in enumerate(detail.columns, 1):
            wd.column_dimensions[get_column_letter(ci)].width = det_widths.get(col, 14)

    return {
        "billed_bal": float(billed["_bal"].sum()),
        "billed_cnt": int(billed["BILL NUMBER"].nunique()),
        "unbilled_bal": float(unbilled["_bal"].sum()),
        "unbilled_cnt": int(unbilled["BILL NUMBER"].nunique()),
        "total_bal": float(base["_bal"].sum()),
    }


# =============================================================================
# EMAIL (classic-Outlook COM, no IT setup required)
# -----------------------------------------------------------------------------
# COM automation is only implemented by classic Outlook - new Outlook
# (olk.exe) has no COM interface - so the sender attaches to a running
# classic Outlook (or starts one itself), then forces a send/receive and
# waits for the Outbox to drain so the message actually transmits even while
# new Outlook is the one open on screen.
# =============================================================================

def _email_body(stats):
    return (
        "Hi Maria,\n\n"
        "Attached is the latest Akzo aging summary.\n\n"
        f"  Billed:   ${stats['billed_bal']:,.2f}  ({stats['billed_cnt']} open invoices)\n"
        f"  Unbilled: ${stats['unbilled_bal']:,.2f}  ({stats['unbilled_cnt']} open invoices)\n"
        f"  Total:    ${stats['total_bal']:,.2f}\n\n"
        "The Overview tab has the headline numbers, the by-Location tabs break "
        "it down, and the Detail tab lists every record on the dashboard "
        "(VLOOKUP on SID or FB#). Figures match the Akzo aging dashboard.\n\n"
        f"Generated {pd.Timestamp.now():%Y-%m-%d %H:%M}.\n\n"
        "Poojan"
    )


def _new_outlook_running():
    """True if 'new' Outlook (olk.exe) is running - it has no COM interface,
    so automation goes through classic outlook.exe regardless."""
    try:
        out = subprocess.run(
            ["tasklist", "/FI", "IMAGENAME eq olk.exe", "/FO", "CSV", "/NH"],
            capture_output=True, text=True, timeout=15).stdout
        return "olk.exe" in out.lower()
    except Exception:
        return False


def _use_new_outlook_flag():
    """HKCU UseNewOutlook=1 (the 'new Outlook' toggle) makes outlook.exe
    itself redirect to olk.exe, which also breaks COM activation
    ('Server execution failed', -2146959355)."""
    try:
        import winreg
        with winreg.OpenKey(
                winreg.HKEY_CURRENT_USER,
                r"Software\Microsoft\Office\16.0\Outlook\Options\General") as k:
            return winreg.QueryValueEx(k, "UseNewOutlook")[0] == 1
    except Exception:
        return False


def _launch_classic_outlook():
    """Shell-launch classic outlook.exe; returns True if a launch started."""
    for p in (
        Path(os.environ.get("ProgramFiles", r"C:\Program Files"))
        / r"Microsoft Office\root\Office16\OUTLOOK.EXE",
        Path(os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)"))
        / r"Microsoft Office\root\Office16\OUTLOOK.EXE",
        Path(os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)"))
        / r"Microsoft Office\Office16\OUTLOOK.EXE",
    ):
        if p.exists():
            subprocess.Popen([str(p)])
            return True
    try:
        os.startfile("outlook.exe")  # resolves via the App Paths registry
        return True
    except OSError:
        return False


def _classic_outlook_app():
    """Attach to classic Outlook, starting it if needed. Plain COM activation
    (Dispatch/DispatchEx) fails with 'Server execution failed' when the new-
    Outlook toggle is in the way, so prefer attaching to a running instance
    and shell-launch classic Outlook ourselves before retrying. Returns
    (app, we_started_it)."""
    import win32com.client as win32
    if _use_new_outlook_flag():
        print("  WARNING: Windows is set to redirect Outlook to the 'new' version "
              "(UseNewOutlook=1). Classic Outlook may refuse to start via "
              "automation; if this send fails, open classic Outlook manually "
              "(Start menu > 'Outlook (classic)') - it can sit in the "
              "background alongside new Outlook - and rerun.")
    launched = False
    last_err = None
    deadline = time.time() + 120
    while time.time() < deadline:
        for attach in (lambda: win32.GetActiveObject("Outlook.Application"),
                       lambda: win32.Dispatch("Outlook.Application")):
            try:
                return attach(), launched
            except Exception as e:
                last_err = e
        if not launched:
            launched = _launch_classic_outlook()
            if not launched:
                break
            print("  Starting classic Outlook (outlook.exe) for the send...")
        time.sleep(3)
    raise RuntimeError(
        f"could not start or attach to classic Outlook for COM automation "
        f"(new Outlook has no COM interface): {last_err}")


def send_via_outlook_com(attachment_path, subject, body, draft_only=False):
    """Classic-Outlook COM send. Attaches to a running classic Outlook or
    starts one, then forces a send/receive and waits for the Outbox to drain
    so the message actually transmits instead of sitting queued until the
    next time classic Outlook happens to be open."""
    if _new_outlook_running():
        print("  Note: 'new' Outlook (olk.exe) is open; it has no COM interface, "
              "so this send needs classic Outlook.")
    outlook, we_started_it = _classic_outlook_app()
    mail = outlook.CreateItem(0)
    mail.To = EMAIL_TO
    mail.Subject = subject
    mail.Body = body
    mail.Attachments.Add(str(attachment_path))
    if draft_only:
        mail.Save()  # lands in classic Drafts even if no window can be shown
        try:
            mail.Display()
        except Exception:
            pass
        return "draft"
    mail.Send()

    ns = outlook.GetNamespace("MAPI")
    outbox = ns.GetDefaultFolder(4)  # olFolderOutbox
    try:
        for sync in ns.SyncObjects:
            sync.Start()
    except Exception:
        pass
    deadline = time.time() + 90
    result = "queued"  # transmits next time classic Outlook is open and online
    while time.time() < deadline:
        if outbox.Items.Count == 0:
            result = "sent"
            break
        time.sleep(2)
    if we_started_it and result == "sent":
        try:
            outlook.Quit()
        except Exception:
            pass
    return result


# =============================================================================
# MAIN
# =============================================================================

def main():
    t0 = time.time()
    print("Akzo Aging Pipeline")
    print(f"  IB 2025: {IB_CAD_2025.exists()} | IB 2026: {IB_CAD_2026.exists()} | "
          f"BU lookup: {BU_LOOKUP_PATH.exists()}")

    credential = InteractiveBrowserCredential()

    soa = load_soa(credential)
    tmw = load_tmw(soa)
    ib_bu_map, ib_stmt_map = load_ib()
    bu_lookup = load_bu_lookup()

    output = build_aging(soa, tmw, ib_bu_map, ib_stmt_map, bu_lookup)

    # --- Working file (internal) ---
    AGING_DIR.mkdir(parents=True, exist_ok=True)
    print(f"\nWriting working file: {OUTPUT_PATH}")
    with pd.ExcelWriter(OUTPUT_PATH, engine="openpyxl") as writer:
        output.to_excel(writer, sheet_name="Aging File", index=False)

    # --- Validation ---
    total_bal = pd.to_numeric(output["BALANCE"], errors="coerce").sum()
    print("\nVALIDATION")
    print(f"  Records: {len(output):,} | "
          f"USD {int((output['Currency'] == 'USD').sum()):,} / "
          f"CAD {int((output['Currency'] == 'CAD').sum()):,}")
    print(f"  Overdue Y {int((output['Overdue'] == 'Y').sum()):,} / "
          f"N {int((output['Overdue'] == 'N').sum()):,}")
    print(f"  null Final BU: {int((output['Final BU'] == 'null').sum())} | "
          f"flagged Remove: {int((output['Remove'] == 'Remove').sum())}")
    print(f"  Total Balance: ${total_bal:,.2f}")

    # --- Maria report ---
    print("\nBuilding Maria report (dashboard-accurate)...")
    stats = build_maria_report(output, MARIA_REPORT_PATH)
    print(f"  Billed:   ${stats['billed_bal']:,.2f} ({stats['billed_cnt']} invoices)")
    print(f"  Unbilled: ${stats['unbilled_bal']:,.2f} ({stats['unbilled_cnt']} invoices)")
    print(f"  Total:    ${stats['total_bal']:,.2f}")
    print(f"  Written:  {MARIA_REPORT_PATH}")

    # --- Email (Maria report only) ---
    subject = f"Akzo Aging Report - {pd.Timestamp.now():%Y-%m-%d}"
    body = _email_body(stats)
    if not EMAIL_ENABLED:
        print("\nEmail skipped (--no-email).")
    elif MARIA_REPORT_PATH.exists() and len(output) > 0:
        try:
            action = send_via_outlook_com(MARIA_REPORT_PATH, subject, body,
                                          draft_only=EMAIL_DRAFT_ONLY)
            if action == "draft":
                print(f"\nDraft created in classic Outlook Drafts - to {EMAIL_TO} "
                      "(a window opens if it can; otherwise find it under Drafts).")
            elif action == "queued":
                print(f"\nEmail QUEUED in classic Outlook's Outbox for {EMAIL_TO} - "
                      "it transmits the next time classic Outlook is open and online.")
            else:
                print(f"\nEmail SENT via classic Outlook to {EMAIL_TO}.")
        except Exception as e:
            print(f"\nEMAIL FAILED (report saved at {MARIA_REPORT_PATH}):\n  {e}")
            print("  Tip: open classic Outlook manually (it can run alongside new "
                  "Outlook) and rerun - the script attaches to the running instance.")
    else:
        print("\nEmail skipped - report missing or output empty.")

    print(f"\nDone in {time.time() - t0:.1f}s ({(time.time() - t0) / 60:.1f} min).")


if __name__ == "__main__":
    main()
