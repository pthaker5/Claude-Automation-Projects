"""
Akzo Nobel Month-Close — BASELINE BUILDER
=========================================
Rebuild every reference baseline FROM RAW QUERY DATA with a single run, so the
baselines no longer have to be hand-built with Excel pivots (the manual pivots
were silently merging / mis-merging case-variant lane keys, which is the main
reason month-close numbers stopped matching).

What this builds
----------------
  1. Expedite Baseline   -> "Baseline Table Expedite Count"
                            (per-lane normal/expedite counts, avg cost,
                             baseline %, cost-increase premium)
  2. LTL RFP Baseline     -> "LTL Baseline File"
                            (per LTL-Bid-ID baseline CPP = cost / weight)
  3. Light-Weight TL->LTL -> "Light Weight TL to LTL Baseline"
                            (per-lane >15k / <15k counts + baseline LT%)
     plus              -> "IN Scope Lanes Light Weight" (carried over / curated)

Refresh with a click
---------------------
    python build_baselines.py                 # rebuild from SQL (production)
    python build_baselines.py --source excel  # rebuild from a raw Excel export
    python build_baselines.py --validate      # rebuild + compare to current sheets

The output is written to the reference workbooks the month-close engine already
reads (akzo_month_close_engine.py), so the engine needs no changes to consume
the refreshed baselines.

=== WHY NUMBERS DRIFTED (root cause) ===
Raw TMS data contains the SAME physical lane spelled in different cases, e.g.
    Wood_Greensboro.NC   and   Wood_GREENSBORO.NC
A hand pivot may or may not merge those; code that groups case-sensitively
splits them into two baseline rows. When the month-close engine later upper-cases
the current-month lane key and left-joins the baseline, two baseline rows match
the same key -> double counting. EVERY grouping key in this builder is
normalized to UPPER-CASE so a lane is one lane, always.
"""

import os
import sys
import math
import argparse
import getpass

import numpy as np
import pandas as pd
import openpyxl


# ===========================================================================
# CONFIG
# ===========================================================================

# --- Baseline windows (bid-event periods). Override per baseline as needed. ---
# Expedite baseline period (the locked window of record):
EXPEDITE_BASELINE_START = "2024-08-01"
EXPEDITE_BASELINE_END   = "2025-08-31"
# Expedite "per month" divisor. Fixed at 12 to match the established methodology
# (per business decision). Set to None to auto-detect from the distinct
# year-months present in the data instead.
EXPEDITE_MONTHS_IN_WINDOW = 12

# LTL RFP + Light-Weight baselines are locked at their respective bid events.
# Point these at the raw query export covering THAT bid analysis window.
LTL_BASELINE_START = "2024-08-01"
LTL_BASELINE_END   = "2025-08-31"
LW_BASELINE_START  = "2024-08-01"
LW_BASELINE_END    = "2025-08-31"

LW_WEIGHT_THRESHOLD = 15000  # lbs; >threshold = heavy band, <=threshold = light

# --- Output reference workbooks (what the month-close engine reads) ---
REF_DIR           = os.path.join(os.path.dirname(os.path.abspath(__file__)), "reference_baselines")
EXPEDITE_BASELINE = os.path.join(REF_DIR, "EXPEDITE REDUCTION Aug 2024 - Aug 2025 BASELINE.xlsx")
REFERENCE_712     = os.path.join(REF_DIR, "712 For Month Close.xlsx")

EXPEDITE_SHEET = "Baseline Table Expedite Count"
LTL_SHEET      = "LTL Baseline File"
LW_SHEET       = "Light Weight TL to LTL Baseline"
INSCOPE_SHEET  = "IN Scope Lanes Light Weight"

# --- Raw source for offline / validation runs (--source excel) ---
# A workbook + sheet holding the raw shipment query (one row per SID).
RAW_EXCEL_PATH  = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                               "sample_data",
                               "EXPEDITE REDUCTION Aug 2024 - Aug 2025 BASELINE_sample.xlsx")
RAW_EXCEL_SHEET = "2025 Query Baseline"

# --- DB (same source as the month-close engine) ---
SERVER   = "az-bwprod.chemlogix.com"
DATABASE = "CLXDW"
DRIVER   = "ODBC Driver 18 for SQL Server"

# Raw query: one row per shipment over the baseline window. Mirrors the columns
# the manual "Query Baseline" tab carried so every baseline can be derived.
BASELINE_QUERY = """
SELECT
    SID, BU, Lane,
    [Origin Loc Code], [Origin City], [Or State], [Origin Zip], [Origin Country],
    [Dest Loc Code],   [Destination City], [Dest State], [Dest Zip], [Dest Country],
    [Pick Up Date], [Movement Type], [Updated Movement Type], [Transport Mode],
    [Priority], [Normalized Weight], [Normalized Ship't Actual Cost],
    [Normalized Adj LineHaul], [Normalized Fuel Charges], [Normalized Base Charges]
FROM dbo.TMSCL712_3_FI_Billing_Extract_Akzo
WHERE [Pick Up Date] >= '{start}' AND [Pick Up Date] <= '{end}'
""".strip()


# ===========================================================================
# BU NORMALIZATION  (kept identical to akzo_month_close_engine.py)
# ===========================================================================

BU_NORMALIZE = {
    "MPY": "MPY", "M & PC": "MPY", "M&PC": "MPY",
    "VR/ SPECIALTY": "VR/ Specialty", "VR/SPECIALTY": "VR/ Specialty",
    "VR/ Specialty": "VR/ Specialty",
    "METAL": "Metal", "Metal": "Metal",
    "WOOD": "Wood", "Wood": "Wood",
    "POWDER": "Powder", "Powder": "Powder",
}


def normalize_bu(raw_bu):
    if raw_bu is None or (isinstance(raw_bu, float) and pd.isna(raw_bu)):
        return None
    return BU_NORMALIZE.get(str(raw_bu).strip(), str(raw_bu).strip())


# ===========================================================================
# HELPERS
# ===========================================================================

def _num(series):
    return pd.to_numeric(series, errors="coerce")


def _is_expedite(priority_series):
    """Expedite = Priority starts with 'EXP' (case-insensitive)."""
    return priority_series.astype(str).str.upper().str.strip().str.startswith("EXP")


def _country_code(country_raw, zip_val):
    """USA / CAN, matching the month-close engine's derivation."""
    c = str(country_raw).strip().upper()
    if c in ("USA", "US", "UNITED STATES"):
        return "USA"
    if c in ("CAN", "CA", "CANADA"):
        return "CAN"
    if zip_val and any(ch.isalpha() for ch in str(zip_val)):
        return "CAN"
    return "USA"


def _norm_lane_from_components(df):
    """Build the canonical UPPER-CASE lane key BU_OriginCity.OrState.

    Prefers an existing 'Lane' column (already present in the TMS query) and
    only rebuilds from components when it is missing. Always upper-cased so
    case variants collapse to one lane.
    """
    if "Lane" in df.columns and df["Lane"].notna().any():
        lane = df["Lane"].astype(str)
    else:
        bu = df["BU"].apply(normalize_bu).astype(str)
        lane = (bu + "_" + df["Origin City"].astype(str).str.strip()
                + "." + df["Or State"].astype(str).str.strip())
    return lane.str.strip().str.upper()


# ===========================================================================
# RAW DATA LOADING
# ===========================================================================

def load_raw_from_sql(start, end, uid, pwd):
    import pyodbc
    conn_str = (
        f"DRIVER={{{DRIVER}}};SERVER={SERVER};DATABASE={DATABASE};"
        f"UID={uid};PWD={pwd};TrustServerCertificate=yes;Encrypt=yes;"
    )
    conn = pyodbc.connect(conn_str)
    try:
        df = pd.read_sql(BASELINE_QUERY.format(start=start, end=end), conn)
    finally:
        conn.close()
    return df


def load_raw_from_excel(path=RAW_EXCEL_PATH, sheet=RAW_EXCEL_SHEET):
    if not os.path.exists(path):
        raise FileNotFoundError(f"Raw Excel source not found: {path}")
    return pd.read_excel(path, sheet_name=sheet, engine="openpyxl")


# ===========================================================================
# BASELINE 1 — EXPEDITE
# ===========================================================================

def build_expedite_baseline(raw, months_in_window=EXPEDITE_MONTHS_IN_WINDOW,
                            scope_lanes=None):
    """Reproduce 'Baseline Table Expedite Count' from raw shipments.

    Filters out PARCEL, groups by UPPER-CASE lane, splits Normal vs Expedite.
    Output columns match what akzo_month_close_engine.load_expedite_baseline reads:
        BU, Origin City, Or State, Lane, Total SID, Normal SID, Normal Cost,
        Expedite Count, Expedite Cost, Baseline %, Cost Increase

    months_in_window: divisor turning the window's total expedite count into a
    per-month rate. None -> auto-detect from the distinct year-months present in
    the data (Aug'24-Aug'25 -> 13). Pass an int to force a fixed divisor.

    scope_lanes: optional iterable of UPPER-CASE lane keys. The expedite project
    tracks a fixed, curated set of in-scope lanes (155 in the original manual
    table — 29 with expedites + 126 kept at zero). When provided, the output is
    restricted to those lanes so a rebuild reproduces the manual table exactly
    instead of expanding to every lane that ever shipped. When omitted, the full
    lane universe is emitted (with a note) — extra zero-baseline lanes are inert
    in the month-close engine (their savings sign-check to $0).
    """
    df = raw.copy()
    df["_tm"]   = df["Transport Mode"].astype(str).str.upper()
    df["_cost"] = _num(df["Normalized Ship't Actual Cost"])
    df["_exp"]  = _is_expedite(df["Priority"])
    df["_lane"] = _norm_lane_from_components(df)

    # Auto-detect the number of months represented (distinct year-months)
    if months_in_window is None:
        if "Pick Up Date" in df.columns:
            dts = pd.to_datetime(df["Pick Up Date"], errors="coerce").dropna()
            months_in_window = dts.dt.to_period("M").nunique() if len(dts) else 12
        else:
            months_in_window = 12
        print(f"  [expedite] months in window auto-detected = {months_in_window}")
    months_in_window = max(int(months_in_window), 1)

    # Parcel never counts toward expedite or normal averages (matches manual pivot)
    df = df[~df["_tm"].str.contains("PARCEL", na=False)]

    # Representative BU / city / state per lane (first non-null), for readability
    meta = (df.assign(_bu=df["BU"].apply(normalize_bu))
              .groupby("_lane")
              .agg(BU=("_bu", "first"),
                   **{"Origin City": ("Origin City", "first"),
                      "Or State":    ("Or State", "first")}))

    grp = df.groupby("_lane")
    norm = df[~df["_exp"]].groupby("_lane")
    expd = df[df["_exp"]].groupby("_lane")

    out = pd.DataFrame({
        "Total SID":      grp.size(),
        "Normal SID":     norm.size(),
        "Normal Cost":    norm["_cost"].mean(),
        "Expedite Count": expd.size(),
        "Expedite Cost":  expd["_cost"].mean(),
    })
    out = meta.join(out, how="right")
    out["Normal SID"]     = out["Normal SID"].fillna(0).astype(int)
    out["Expedite Count"] = out["Expedite Count"].fillna(0).astype(int)
    out["Total SID"]      = out["Total SID"].fillna(0).astype(int)

    # Baseline % = expedite count / normal count ; Cost Increase = exp/normal - 1
    out["Baseline %"] = np.where(out["Normal SID"] > 0,
                                 out["Expedite Count"] / out["Normal SID"], 0.0)
    out["Cost Increase"] = np.where(
        (out["Normal Cost"] > 0) & out["Expedite Cost"].notna(),
        out["Expedite Cost"] / out["Normal Cost"] - 1.0, np.nan)

    # Derived field the engine recomputes anyway, included for transparency
    out["baseline_avg_per_month"] = np.ceil(out["Expedite Count"] / months_in_window).astype(int)

    out = out.reset_index().rename(columns={"_lane": "Lane"})

    if scope_lanes is not None:
        scope = {str(l).strip().upper() for l in scope_lanes}
        out = out[out["Lane"].isin(scope)]
        missing = scope - set(out["Lane"])
        if missing:
            print(f"  [expedite] WARNING: {len(missing)} in-scope lanes had no "
                  f"shipments in the window: {sorted(missing)[:5]}{' ...' if len(missing) > 5 else ''}")

    cols = ["BU", "Origin City", "Or State", "Lane", "Total SID", "Normal SID",
            "Normal Cost", "Expedite Count", "Expedite Cost", "Baseline %",
            "Cost Increase", "baseline_avg_per_month"]
    return out[cols].sort_values("Expedite Count", ascending=False).reset_index(drop=True)


# ===========================================================================
# BASELINE 2 — LTL RFP
# ===========================================================================

def build_ltl_baseline(raw):
    """Reproduce 'LTL Baseline File': baseline CPP per LTL Bid ID.

    Bid ID = BU_OriginZip.OriginCountry_DestZip.DestCountry  (UPPER-CASE)
    BASELINE CPP = sum(cost) / sum(weight)
    Output columns match akzo_month_close_engine.load_ltl_lane_baselines:
        Row Labels, Sum of Normalized Ship't Actual Cost,
        Sum of Normalized Weight, BASELINE CPP, BASELINE SID
    """
    df = raw.copy()
    df["_tm"] = df["Transport Mode"].astype(str).str.upper()
    df = df[df["_tm"].str.contains("LTL", na=False)]

    df["_cost"] = _num(df["Normalized Ship't Actual Cost"])
    df["_wt"]   = _num(df["Normalized Weight"])
    df["_bu"]   = df["BU"].apply(normalize_bu)
    df["_oc"]   = df.apply(lambda r: _country_code(r.get("Origin Country", ""), r.get("Origin Zip", "")), axis=1)
    df["_dc"]   = df.apply(lambda r: _country_code(r.get("Dest Country", ""), r.get("Dest Zip", "")), axis=1)

    df["_bid"] = (
        df["_bu"].astype(str).str.strip() + "_"
        + df["Origin Zip"].astype(str).str.strip() + "." + df["_oc"] + "_"
        + df["Dest Zip"].astype(str).str.strip() + "." + df["_dc"]
    ).str.upper()

    df = df[df["_cost"].notna() & (df["_cost"] > 0) & df["_wt"].notna() & (df["_wt"] > 0)]

    g = df.groupby("_bid").agg(
        cost=("_cost", "sum"),
        wt=("_wt", "sum"),
        sid=("SID", "count"),
    )
    g["cpp"] = g["cost"] / g["wt"]
    g = g.reset_index().rename(columns={
        "_bid": "Row Labels",
        "cost": "Sum of Normalized Ship't Actual Cost",
        "wt":   "Sum of Normalized Weight",
        "cpp":  "BASELINE CPP",
        "sid":  "BASELINE SID",
    })
    return g.sort_values("Sum of Normalized Ship't Actual Cost", ascending=False).reset_index(drop=True)


# ===========================================================================
# BASELINE 3 — LIGHT-WEIGHT TL -> LTL
# ===========================================================================

def build_lw_baseline(raw, threshold=LW_WEIGHT_THRESHOLD):
    """Reproduce 'Light Weight TL to LTL Baseline': per-lane heavy/light counts + LT%.

    Truckload only, grouped by UPPER-CASE lane.
    Output columns match akzo_month_close_engine.load_lw_baseline (col index 4 = LT%):
        Row Labels, Greater than 15000, Less than 15000, Grand Total, Column1
    """
    df = raw.copy()
    df["_tm"] = df["Transport Mode"].astype(str).str.upper()
    df = df[df["_tm"] == "TRUCKLOAD"]

    df["_wt"]   = _num(df["Normalized Weight"])
    df["_lane"] = _norm_lane_from_components(df)
    df["_heavy"] = df["_wt"] > threshold

    g = df.groupby("_lane")
    out = pd.DataFrame({
        "Greater than 15000": g.apply(lambda x: int((x["_wt"] > threshold).sum()), include_groups=False),
        "Less than 15000":    g.apply(lambda x: int((x["_wt"] <= threshold).sum()), include_groups=False),
    })
    out["Grand Total"] = out["Greater than 15000"] + out["Less than 15000"]
    out["Column1"] = np.where(out["Grand Total"] > 0,
                              out["Less than 15000"] / out["Grand Total"], 0.0)
    out = out.reset_index().rename(columns={"_lane": "Row Labels"})
    return out.sort_values("Grand Total", ascending=False).reset_index(drop=True)


def carry_over_expedite_scope(expedite_path=EXPEDITE_BASELINE, sheet=EXPEDITE_SHEET):
    """The expedite project tracks a fixed, curated set of in-scope lanes. Read
    them (UPPER-CASE) from the existing baseline workbook so a rebuild keeps the
    same scope. Returns a set, or None if the file/sheet is unavailable.
    """
    if not os.path.exists(expedite_path):
        return None
    try:
        cur = pd.read_excel(expedite_path, sheet_name=sheet, engine="openpyxl")
    except Exception:
        return None
    if "Lane" not in cur.columns:
        return None
    return {str(l).strip().upper() for l in cur["Lane"].dropna()}


def carry_over_in_scope_lanes(reference_712_path=REFERENCE_712, sheet=INSCOPE_SHEET):
    """In-scope LW lanes are a curated project-scope selection (set at the bid
    event), NOT auto-derived. Carry them over from the existing reference file
    so a rebuild never silently changes which 36 lanes are in scope.
    Returns a list (UPPER-CASE) or [] if the file/sheet is unavailable.
    """
    if not os.path.exists(reference_712_path):
        return []
    wb = openpyxl.load_workbook(reference_712_path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        wb.close()
        return []
    ws = wb[sheet]
    lanes = []
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        if row[0]:
            lanes.append(str(row[0]).strip().upper())
    wb.close()
    return lanes


# ===========================================================================
# WRITERS  (replace a single sheet, preserve everything else in the workbook)
# ===========================================================================

def write_sheet(workbook_path, sheet_name, df, header_extra=None):
    """Overwrite one sheet of an existing workbook with df, keeping other sheets.

    Creates the workbook if it does not exist. header_extra is an optional list
    of cell values written above the table header (pivot-style banner rows).
    """
    if os.path.exists(workbook_path):
        wb = openpyxl.load_workbook(workbook_path)
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
        ws = wb.create_sheet(sheet_name)
        # move new sheet to front for visibility
        wb.move_sheet(sheet_name, -(len(wb.sheetnames) - 1))
    else:
        os.makedirs(os.path.dirname(workbook_path), exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name

    r = 1
    if header_extra:
        for line in header_extra:
            ws.cell(row=r, column=1).value = line
            r += 1
    # header
    for c, col in enumerate(df.columns, start=1):
        ws.cell(row=r, column=c).value = str(col)
    r += 1
    for _, row in df.iterrows():
        for c, val in enumerate(row.tolist(), start=1):
            if isinstance(val, (np.integer,)):
                val = int(val)
            elif isinstance(val, (np.floating,)):
                val = float(val)
            ws.cell(row=r, column=c).value = val
        r += 1

    wb.save(workbook_path)
    wb.close()


def write_inscope_lanes(workbook_path, lanes, sheet_name=INSCOPE_SHEET):
    df = pd.DataFrame({"IN Scope Lanes Light Weight": lanes})
    write_sheet(workbook_path, sheet_name, df)


# ===========================================================================
# VALIDATION  (compare a freshly built table to the one currently in the file)
# ===========================================================================

def _read_existing(workbook_path, sheet_name, skiprows=0):
    if not os.path.exists(workbook_path):
        return None
    try:
        return pd.read_excel(workbook_path, sheet_name=sheet_name,
                             skiprows=skiprows, engine="openpyxl")
    except Exception:
        return None


def validate_expedite(built, existing_path, sheet=EXPEDITE_SHEET, tol=0.01):
    cur = _read_existing(existing_path, sheet)
    if cur is None:
        print("  [validate] no existing expedite sheet to compare")
        return
    cur = cur.copy()
    cur["_lane"] = cur["Lane"].astype(str).str.upper()
    b = built.set_index(built["Lane"].astype(str).str.upper())
    c = cur.set_index("_lane")
    common = b.index.intersection(c.index)
    mism = 0
    for lane in common:
        for col in ["Expedite Count", "Normal SID", "Total SID"]:
            bv = float(b.loc[lane, col]); cv = float(c.loc[lane, col])
            if abs(bv - cv) > tol:
                mism += 1
                if mism <= 10:
                    print(f"    DIFF {lane} {col}: built={bv} existing={cv}")
    only_built = b.index.difference(c.index)
    only_cur   = c.index.difference(b.index)
    print(f"  [validate expedite] lanes built={len(b)} existing={len(c)} "
          f"common={len(common)} cell-mismatches={mism} "
          f"only-in-built={len(only_built)} only-in-existing={len(only_cur)}")


# ===========================================================================
# MAIN
# ===========================================================================

def _load_raw(source):
    if source == "excel":
        print(f"Loading raw from Excel: {RAW_EXCEL_PATH} :: {RAW_EXCEL_SHEET}")
        raw = load_raw_from_excel()  # one window/file feeds all three baselines
        return {"expedite": raw, "ltl": raw, "lw": raw}
    # SQL: each baseline may use a different bid window
    uid = input("SQL Username: ").strip()
    pwd = getpass.getpass("SQL Password: ")
    print(f"Pull expedite window {EXPEDITE_BASELINE_START}..{EXPEDITE_BASELINE_END}")
    exp = load_raw_from_sql(EXPEDITE_BASELINE_START, EXPEDITE_BASELINE_END, uid, pwd)
    print(f"Pull LTL window      {LTL_BASELINE_START}..{LTL_BASELINE_END}")
    ltl = load_raw_from_sql(LTL_BASELINE_START, LTL_BASELINE_END, uid, pwd)
    print(f"Pull LW window       {LW_BASELINE_START}..{LW_BASELINE_END}")
    lw  = load_raw_from_sql(LW_BASELINE_START, LW_BASELINE_END, uid, pwd)
    return {"expedite": exp, "ltl": ltl, "lw": lw}


def main():
    ap = argparse.ArgumentParser(description="Rebuild Akzo month-close baselines from raw data.")
    ap.add_argument("--source", choices=["sql", "excel"], default="sql",
                    help="raw data source (default: sql)")
    ap.add_argument("--validate", action="store_true",
                    help="compare freshly built tables to the current reference sheets")
    ap.add_argument("--no-write", action="store_true",
                    help="build + report only; do not write workbooks")
    args = ap.parse_args()

    print("=" * 70)
    print("Akzo Month-Close — Baseline Builder")
    print("=" * 70)

    raw = _load_raw(args.source)

    print("\n[1/3] Expedite baseline ...")
    exp_scope = carry_over_expedite_scope()
    if exp_scope:
        print(f"  carrying over {len(exp_scope)} in-scope expedite lanes")
    else:
        print("  no existing expedite scope found -> emitting full lane universe")
    exp_base = build_expedite_baseline(raw["expedite"], scope_lanes=exp_scope)
    print(f"  {len(exp_base)} lanes  |  total expedite SIDs = {int(exp_base['Expedite Count'].sum()):,}")

    print("[2/3] LTL RFP baseline ...")
    ltl_base = build_ltl_baseline(raw["ltl"])
    print(f"  {len(ltl_base)} bid lanes  |  avg CPP = {ltl_base['BASELINE CPP'].mean():.4f}")

    print("[3/3] Light-Weight baseline ...")
    lw_base = build_lw_baseline(raw["lw"])
    inscope = carry_over_in_scope_lanes()
    print(f"  {len(lw_base)} lanes  |  in-scope lanes carried over = {len(inscope)}")

    if args.validate:
        print("\n--- Validation against current reference sheets ---")
        validate_expedite(exp_base, EXPEDITE_BASELINE)

    if args.no_write:
        print("\n--no-write set: not writing workbooks. Done.")
        return

    print("\nWriting reference workbooks ...")
    write_sheet(EXPEDITE_BASELINE, EXPEDITE_SHEET, exp_base)
    print(f"  [OK] {EXPEDITE_SHEET} -> {EXPEDITE_BASELINE}")
    write_sheet(REFERENCE_712, LTL_SHEET, ltl_base)
    print(f"  [OK] {LTL_SHEET} -> {REFERENCE_712}")
    write_sheet(REFERENCE_712, LW_SHEET, lw_base)
    print(f"  [OK] {LW_SHEET}")
    if inscope:
        write_inscope_lanes(REFERENCE_712, inscope)
        print(f"  [OK] {INSCOPE_SHEET} ({len(inscope)} lanes)")
    print("\nDone. Baselines refreshed.")


if __name__ == "__main__":
    main()
