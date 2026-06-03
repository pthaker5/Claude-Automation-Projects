#!/usr/bin/env python3
"""
edw_pull.py — Pull DC&E billing audit data from Quantix EDW (Azure SQL)
and produce Excel/CSV files matching what billing_audit_v18.html expects.

Usage:
    python edw_pull.py --start 2026-05-11 --end 2026-05-18 [--output-dir ./output]

Outputs (file names must match dashboard exactly):
    HJ - HighJump Revenue by Chargeback.xlsx   (Report 1: Invoice)
    HJ - Lost Revenue.xlsx                     (Report 2: multi-sheet)
    WMS Rates From EDW(CORE core_rates).csv    (Report 3: rate table)
    HJ - Railcar Arrival Report.xlsx           (Report 5: railcar)
    HJ - Revenue Audit Packaging.xlsx          (Report 6: packaging)

Report 4 (No Charges) is TODO — see comment in pull_no_charges().

Dependencies:
    pip install pyodbc pandas openpyxl azure-identity
"""

import argparse
import os
import struct
import sys
from datetime import datetime
from decimal import Decimal

import pandas as pd
import pyodbc
from azure.identity import ClientSecretCredential
from openpyxl import Workbook

# ── Connection settings ───────────────────────────────────────────────────────
# Service-principal credentials come from the environment (see .env.example /
# server/sql_client.py). Never hardcode the client secret here.
SERVER        = os.getenv("EDW_DEV_SERVER", "dev-quantix-useast-sql.database.windows.net")
DATABASE      = os.getenv("EDW_DEV_DB", "EDW_Dev")
TENANT        = os.getenv("EDW_SP_TENANT_ID", "")
CLIENT_ID     = os.getenv("EDW_SP_CLIENT_ID", "")
CLIENT_SECRET = os.getenv("EDW_SP_CLIENT_SECRET", "")
SCOPE         = "https://database.windows.net/.default"

if not (TENANT and CLIENT_ID and CLIENT_SECRET):
    print("Set EDW_SP_TENANT_ID, EDW_SP_CLIENT_ID, EDW_SP_CLIENT_SECRET in the environment.", file=sys.stderr)

# Batch size for IN-clause queries (SQL Server limit is 2100 params per stmt)
BATCH_SIZE = 500


# ── Connection ────────────────────────────────────────────────────────────────

def get_connection():
    """Azure SQL connection via MSAL service-principal token auth."""
    credential = ClientSecretCredential(TENANT, CLIENT_ID, CLIENT_SECRET)
    token = credential.get_token(SCOPE)
    token_bytes = token.token.encode("utf-16-le")
    token_struct = struct.pack(f"<I{len(token_bytes)}s", len(token_bytes), token_bytes)
    conn_str = (
        "Driver={ODBC Driver 18 for SQL Server};"
        f"Server={SERVER};Database={DATABASE};"
        "Encrypt=yes;TrustServerCertificate=yes;"
    )
    conn = pyodbc.connect(conn_str, attrs_before={1256: token_struct}, timeout=60)
    conn.timeout = 300
    return conn


# ── Helpers ───────────────────────────────────────────────────────────────────

def _fix_decimals(df):
    """Cast Decimal columns to float (pyodbc returns Decimal for SQL NUMERIC/DECIMAL)."""
    for col in df.columns:
        if not df[col].empty and df[col].apply(lambda x: isinstance(x, Decimal)).any():
            df[col] = df[col].apply(lambda x: float(x) if isinstance(x, Decimal) else x)
    return df


def _run_query(conn, sql, params=None):
    """Execute SQL and return a DataFrame."""
    cursor = conn.cursor()
    if params:
        cursor.execute(sql, params)
    else:
        cursor.execute(sql)
    cols = [c[0] for c in cursor.description]
    rows = cursor.fetchall()
    df = pd.DataFrame.from_records(rows, columns=cols)
    return _fix_decimals(df)


def _query_in_batches(conn, sql_template, ids, batch_size=BATCH_SIZE):
    """
    Execute a query with an IN clause, batching to avoid pyodbc parameter limits.

    sql_template must contain exactly one `{}` where the comma-separated `?`
    placeholders will be inserted.  Example:
        "SELECT * FROM t WHERE order_number IN ({})"
    """
    if not ids:
        return pd.DataFrame()

    ids = list(dict.fromkeys(ids))  # deduplicate, preserve order
    chunks = [ids[i : i + batch_size] for i in range(0, len(ids), batch_size)]
    dfs = []
    for chunk in chunks:
        placeholders = ",".join("?" * len(chunk))
        sql = sql_template.format(placeholders)
        cursor = conn.cursor()
        cursor.execute(sql, chunk)
        cols = [c[0] for c in cursor.description]
        rows = cursor.fetchall()
        df = pd.DataFrame.from_records(rows, columns=cols)
        dfs.append(_fix_decimals(df))

    return pd.concat(dfs, ignore_index=True) if dfs else pd.DataFrame()


# ── Report 1: Invoice — Revenue by Chargeback ─────────────────────────────────

def pull_invoice(conn, start, end, output_dir):
    """
    Pulls t_bmm_charge + billing tables for the date range.
    Output columns match the 14 columns billing_audit_v18.html expects via px().
    Validated: 5/11–5/18 → 9,392 rows, $2,390,495.42.
    """
    print("  Querying Invoice (Revenue by Chargeback)...", flush=True)
    sql = """
SELECT
    cm.wh_id                        AS [WH ID],
    cust.customer_code              AS [Customer Code],
    cust.customer_name              AS [Customer Name],
    cb.chargeback_code              AS [Chargeback Code],
    cb.description                  AS [Charge Description],
    c.rate                          AS [Rate],
    cr.per_text                     AS [Per],
    CAST(c.charge_amount AS FLOAT)  AS [Charge Amount],
    c.charge_date_time              AS [Charge Date],
    c.order_num_value               AS [Order Number],
    inv.invoice_number              AS [Invoice Number],
    inv.closed_date                 AS [Invoice Date],
    c.reason_text_value             AS [Description 2],
    c.prompt_text_value             AS [Qty]
FROM Korber.t_bmm_charge c
JOIN Korber.t_bmm_cont_inv_type_chargeback cb
    ON c.chargeback_id = cb.chargeback_id
JOIN Korber.t_bmm_invoice inv
    ON c.invoice_id = inv.invoice_id
JOIN Korber.t_bmm_contract_invoice_type cit
    ON cb.contract_invoice_type_id = cit.contract_invoice_type_id
JOIN Korber.t_bmm_contract_master cm
    ON cit.contract_id = cm.contract_id
JOIN Korber.t_bmm_customer cust
    ON cm.customer_id = cust.customer_id
LEFT JOIN Korber.t_bmm_chargeback_rate cr
    ON c.chargeback_rate_id = cr.chargeback_rate_id
WHERE inv.closed_date >= ?
  AND inv.closed_date <  ?
  AND c.charge_amount <> 0
"""
    df = _run_query(conn, sql, (start, end))
    path = os.path.join(output_dir, "HJ - HighJump Revenue by Chargeback.xlsx")
    df.to_excel(path, index=False)
    total = df["Charge Amount"].sum() if "Charge Amount" in df.columns else 0
    print(f"  Invoice: {len(df):,} rows, ${total:,.2f}", flush=True)
    return df


# ── Report 3: WMS Rate Table ──────────────────────────────────────────────────

def pull_rate_table(conn, output_dir):
    """
    Pulls CORE.Core_Rates (pre-joined rate table).
    Output is CSV — the dashboard pcsv() parser accepts CSV for rates.
    """
    print("  Querying WMS Rate Table...", flush=True)
    sql = """
SELECT
    customer_code,
    contract_wh_id,
    chargeback_code,
    description,
    rate,
    per_text,
    order_type,
    container_type,
    precedence
FROM CORE.Core_Rates
WHERE source_table IN ('WMS_Rates', 'WMS_Manual')
"""
    df = _run_query(conn, sql)
    path = os.path.join(output_dir, "WMS Rates From EDW(CORE core_rates).csv")
    df.to_csv(path, index=False)
    print(f"  Rate Table: {len(df):,} rows", flush=True)
    return df


# ── Report 2: Lost Revenue (multi-sheet) ──────────────────────────────────────
#
# Format required by pxAllSheets() in billing_audit_v18.html:
#
#   Sheet1  row 0:  canonical column headers  ['wh id', 'WH ID', 'Name', ...]
#           row 1:  section marker            ['43-531', '', '', ...]
#           row 2+: data rows                 [None, '43-531', 'Vinmar...', ...]
#
#   Sheet2+ row 0:  section marker            ['43-533', '', '', ...]
#           row 1+: data rows                 [None, '43-533', ...]
#
# The parser skips any row where col[0] is truthy (section markers and the
# header row itself).  Data rows are identified by col[0] == null/None.
# The canonical header is extracted once from Sheet1 row 0 and used to map
# all data rows from all sheets.
#
# Row granularity: one row per t_bmm_event_log event per order (matching SSRS).
# Each event provides one Comment value.  Billing Link and Our Supplies are
# aggregate counts on the order (same value repeated on every row for that order).

# These 18 columns are the "logical" data columns (after the hidden 'wh id' col 0).
_LR_COLS = [
    "WH ID", "Name", "Client Code", "Order Number", "Amount Billed",
    "Ship Date", "BOL Number", "Order Type", "Carrier Name",
    "Comment", "Source Vessel", "Dest Vessel", "Our Supplies", "Billing Link",
    "Charge", "Corrected Amount", "Difference", "Reason",
]
# Full header row written to Sheet1 row 0: 'wh id' at col 0, data cols follow.
_LR_HEADER = ["wh id"] + _LR_COLS  # 19 columns


def pull_lost_revenue(conn, start, end, output_dir):
    """
    Lost Revenue report: orders shipped in range with billing event linkage.
    Produces a multi-sheet Excel matching pxAllSheets() expected format.
    """
    print("  Querying Lost Revenue orders...", flush=True)

    # ── Step 1: Orders in date range ─────────────────────────────────────────
    orders_df = _run_query(conn, """
SELECT
    o.wh_id                  AS [WH ID],
    o.display_order_number   AS [Order Number],
    cl.name                  AS [Name],
    o.client_code            AS [Client Code],
    o.actual_ship_date       AS [Ship Date],
    o.bol_number             AS [BOL Number],
    o.carrier                AS [Carrier Name],
    o.order_id               AS _order_id
FROM Korber.t_order o
LEFT JOIN Korber.t_client cl
    ON o.client_code = cl.client_code AND o.wh_id = cl.wh_id
WHERE o.actual_ship_date >= ?
  AND o.actual_ship_date <  ?
  AND o.status IN ('SHIPPED', 'COMPLETE')
""", (start, end))

    print(f"    Orders: {len(orders_df):,}", flush=True)
    if orders_df.empty:
        print("  Lost Revenue: 0 rows", flush=True)
        return

    order_nums  = orders_df["Order Number"].dropna().unique().tolist()
    order_ids   = orders_df["_order_id"].dropna().astype(int).unique().tolist()

    # ── Step 2: Vessel info (source/dest) from order_detail ──────────────────
    print("  Querying order detail (vessel info)...", flush=True)
    try:
        vessel_df = _query_in_batches(conn, """
SELECT order_id AS _order_id,
       MIN(source_vessel) AS [Source Vessel],
       MIN(dest_vessel)   AS [Dest Vessel]
FROM Korber.t_order_detail
WHERE order_id IN ({})
GROUP BY order_id
""", order_ids)
    except Exception as e:
        print(f"    Warning: vessel query failed ({e})", flush=True)
        vessel_df = pd.DataFrame(columns=["_order_id", "Source Vessel", "Dest Vessel"])

    # ── Step 3: Order Type lookup ─────────────────────────────────────────────
    print("  Querying order types...", flush=True)
    try:
        otype_df = _query_in_batches(conn, """
SELECT o.order_id AS _order_id, ot.description AS [Order Type]
FROM Korber.t_order o
LEFT JOIN Korber.t_order_type ot ON o.type_id = ot.type_id
WHERE o.order_id IN ({})
""", order_ids)
    except Exception as e:
        # t_order_type might not exist; derive from order number prefix
        print(f"    Warning: order type table query failed ({e}), deriving from prefix", flush=True)
        otype_df = orders_df[["_order_id", "Order Number"]].copy()
        def _guess_type(onum):
            if not onum:
                return None
            p = str(onum).lstrip("0123456789-")[:2].upper()
            return {
                "SO": "Shipping Order", "PW": "Packaging Work Order",
                "RS": "Reclass Order",  "RO": "Receiving Order",
                "WO": "Work Order",
            }.get(p)
        otype_df["Order Type"] = otype_df["Order Number"].apply(_guess_type)
        otype_df = otype_df[["_order_id", "Order Type"]]

    # ── Step 4: Event log — one row per event (Comment + tran_type counts) ───
    print("  Querying event log (billing events)...", flush=True)
    el_df = pd.DataFrame()
    # Try column names in order of likelihood
    for comment_col in ("tran_text", "note_text", "description", "comment_text", "reason"):
        try:
            el_df = _query_in_batches(conn, f"""
SELECT order_number AS [Order Number],
       tran_type,
       {comment_col} AS [Comment]
FROM Korber.t_bmm_event_log
WHERE order_number IN ({{}})
""", order_nums)
            break  # succeeded
        except Exception:
            continue

    if el_df.empty:
        # Fall back: get events without a comment column (correct row count, no comment text)
        try:
            el_df = _query_in_batches(conn, """
SELECT order_number AS [Order Number],
       tran_type,
       CAST(NULL AS NVARCHAR(500)) AS [Comment]
FROM Korber.t_bmm_event_log
WHERE order_number IN ({})
""", order_nums)
        except Exception as e:
            print(f"    Warning: event log query failed entirely ({e}) — falling back to 1 row/order", flush=True)

    # Build per-order aggregates from event log
    billing_counts  = {}  # order_number -> count of tran_type='340'
    supplies_counts = {}  # order_number -> count of tran_type='344'
    if not el_df.empty:
        for _, row in el_df.iterrows():
            onum = row["Order Number"]
            tt   = str(row.get("tran_type") or "")
            if tt == "340":
                billing_counts[onum]  = billing_counts.get(onum, 0)  + 1
            elif tt == "344":
                supplies_counts[onum] = supplies_counts.get(onum, 0) + 1

    # ── Step 5: Amount Billed from t_bmm_charge ───────────────────────────────
    print("  Querying charges (Amount Billed)...", flush=True)
    try:
        charge_df = _query_in_batches(conn, """
SELECT order_num_value AS [Order Number],
       SUM(CAST(charge_amount AS FLOAT)) AS [Amount Billed]
FROM Korber.t_bmm_charge
WHERE order_num_value IN ({})
GROUP BY order_num_value
""", order_nums)
        amount_billed = dict(zip(charge_df["Order Number"], charge_df["Amount Billed"])) if not charge_df.empty else {}
    except Exception as e:
        print(f"    Warning: charges query failed ({e})", flush=True)
        amount_billed = {}

    # ── Step 6: Merge auxiliary data into orders_df ───────────────────────────
    merged = orders_df.copy()
    if not vessel_df.empty:
        merged = merged.merge(vessel_df, on="_order_id", how="left")
    else:
        merged["Source Vessel"] = None
        merged["Dest Vessel"]   = None

    if not otype_df.empty:
        merged = merged.merge(otype_df, on="_order_id", how="left")
    else:
        merged["Order Type"] = None

    # ── Step 7: Build final row list ──────────────────────────────────────────
    # Row granularity: one row per event_log entry per order.
    # If no events for an order, fall back to one row with no comment.
    final_rows = []

    if not el_df.empty:
        el_by_order = el_df.groupby("Order Number")
        for _, ord_row in merged.iterrows():
            onum = ord_row["Order Number"]
            base = _make_lr_base(ord_row, onum, amount_billed, billing_counts, supplies_counts)

            if onum in el_by_order.groups:
                for _, ev in el_by_order.get_group(onum).iterrows():
                    row = dict(base)
                    row["Comment"] = ev.get("Comment") or ""
                    final_rows.append(row)
            else:
                base["Comment"] = ""
                final_rows.append(base)
    else:
        # No event log data — one row per order, no comment
        for _, ord_row in merged.iterrows():
            onum = ord_row["Order Number"]
            row  = _make_lr_base(ord_row, onum, amount_billed, billing_counts, supplies_counts)
            row["Comment"] = ""
            final_rows.append(row)

    # ── Step 8: Write multi-sheet Excel ──────────────────────────────────────
    # Group rows by WH ID, sorted alphabetically
    wh_groups: dict[str, list] = {}
    for row in final_rows:
        wh = row.get("WH ID") or ""
        wh_groups.setdefault(wh, []).append(row)

    wb = Workbook()
    wb.remove(wb.active)  # remove default blank sheet

    for sheet_idx, (wh_id, rows) in enumerate(sorted(wh_groups.items())):
        ws = wb.create_sheet(title=f"Sheet{sheet_idx + 1}")

        # Sheet1 gets the canonical header row at row 0 (all other sheets skip it)
        if sheet_idx == 0:
            ws.append(_LR_HEADER)  # row 0: ['wh id', 'WH ID', 'Name', ...]

        # Section marker: WH ID value in col 0, rest empty (skipped by pxAllSheets)
        ws.append([wh_id] + [""] * (len(_LR_HEADER) - 1))

        # Data rows: None in col 0, WH ID in col 1, then remaining cols
        for row in rows:
            ws.append([None] + [row.get(col) for col in _LR_COLS])

    path = os.path.join(output_dir, "HJ - Lost Revenue.xlsx")
    wb.save(path)

    unique_orders = len({r["Order Number"] for r in final_rows})
    print(
        f"  Lost Revenue: {len(final_rows):,} rows across "
        f"{len(wh_groups)} sheets ({unique_orders:,} unique orders)",
        flush=True,
    )


def _make_lr_base(ord_row, onum, amount_billed, billing_counts, supplies_counts):
    """Build a base Lost Revenue row dict (Comment filled in by caller)."""
    return {
        "WH ID":           ord_row["WH ID"],
        "Name":            ord_row.get("Name"),
        "Client Code":     ord_row["Client Code"],
        "Order Number":    onum,
        "Amount Billed":   amount_billed.get(onum, 0.0),
        "Ship Date":       ord_row["Ship Date"],
        "BOL Number":      ord_row.get("BOL Number"),
        "Order Type":      ord_row.get("Order Type"),
        "Carrier Name":    ord_row.get("Carrier Name"),
        "Comment":         "",  # overwritten by caller
        "Source Vessel":   ord_row.get("Source Vessel"),
        "Dest Vessel":     ord_row.get("Dest Vessel"),
        "Our Supplies":    supplies_counts.get(onum, 0),
        "Billing Link":    billing_counts.get(onum, 0),
        "Charge":          "",
        "Corrected Amount": "",
        "Difference":      "",
        "Reason":          "",
    }


# ── Report 4: No Charges (TODO) ───────────────────────────────────────────────

def pull_no_charges(conn, start, end, output_dir):
    """
    TODO: Billable Activity with No Charges.

    Logic: identify t_bmm_event_log events whose tran_type is in the set of
    billable event types (defined in t_bmm_param_wms_event), then find which
    of those have no matching row in t_bmm_charge for the same order.

    This requires row-level validation against the SSRS export to confirm the
    exact filter conditions.  The conceptual query in EDW_TABLE_MAPPING.md is a
    starting point but has not been validated.
    """
    print("  No Charges: TODO (needs t_bmm_param_wms_event validation)", flush=True)


# ── Report 5: Railcar Arrivals ────────────────────────────────────────────────

def pull_railcar(conn, start, end, output_dir):
    """
    ASN records where vessel = 'RAILCAR' for the delivery date range.
    Output has headers at row 0; pxRailcar()'s standard-parse branch handles it.
    """
    print("  Querying Railcar Arrivals...", flush=True)
    sql = """
SELECT
    am.wh_id              AS [WH ID],
    am.client_code        AS [Client Code],
    cl.name               AS [Client Name],
    ad.hu_id              AS [Railcar],
    am.delivery_date      AS [Arrived Date]
FROM Korber.t_asn_master am
JOIN Korber.t_asn_detail ad
    ON am.asn_number = ad.asn_number AND am.wh_id = ad.wh_id
LEFT JOIN Korber.t_client cl
    ON am.client_code = cl.client_code AND am.wh_id = cl.wh_id
WHERE am.delivery_date >= ?
  AND am.delivery_date <  ?
  AND ad.vessel = 'RAILCAR'
"""
    df = _run_query(conn, sql, (start, end))
    path = os.path.join(output_dir, "HJ - Railcar Arrival Report.xlsx")
    df.to_excel(path, index=False)
    print(f"  Railcar: {len(df):,} rows", flush=True)
    return df


# ── Report 6: Revenue Audit Packaging ────────────────────────────────────────
#
# SSRS output format (from HJ - Revenue Audit Packaging.xlsx):
#   Row 0: empty
#   Row 1: title row — col 4 = "Revenue Audit Packaging"
#   Row 2: empty
#   Row 3: column headers  (pxPackaging() uses range:3 to read these)
#   Row 4+: data rows
#
# Column headers (26 total, positions 4 and 7 are None):
#   Warehouse, Order Number, Comments, Store Order Number, [None], Type,
#   Customer PO Number, [None], Carrier, BOL Number, Item Description,
#   Lot Number, Created By, Text String, Source Vessel, Destination Vessel,
#   Quantity, Order Date, Actual Ship Date, Client ID, Invoice Number,
#   Closed Date, Contract Invoice Type Code, Description, TRIM Order, Count Trans

_PKG_HEADERS = [
    "Warehouse", "Order Number", "Comments", "Store Order Number", None,
    "Type", "Customer PO Number", None,
    "Carrier", "BOL Number", "Item Description", "Lot Number", "Created By",
    "Text String", "Source Vessel", "Destination Vessel", "Quantity",
    "Order Date", "Actual Ship Date", "Client ID",
    "Invoice Number", "Closed Date", "Contract Invoice Type Code", "Description",
    "TRIM Order", "Count Trans",
]


def pull_packaging(conn, start, end, output_dir):
    """
    Packaging Work Orders (PW prefix) with billing linkage.
    Reproduces the 4-row preamble + header-at-row-3 format from SSRS.
    pxPackaging() standard-parse branch checks std[0]['Warehouse'] — the
    4-row preamble means standard parse will miss it and fall through to
    the range:3 fallback, which is the correct SSRS-compatible path.
    """
    print("  Querying Revenue Audit Packaging...", flush=True)

    # Main query — best-effort for optional fields that may need schema tuning
    sql = """
SELECT
    o.wh_id                  AS [Warehouse],
    o.display_order_number   AS [Order Number],
    ''                       AS [Comments],  -- comment source TBD; not on t_order or t_order_detail
    o.store_order_number     AS [Store Order Number],
    'Packaging Work Order'   AS [Type],
    o.cust_po_number         AS [Customer PO Number],
    o.carrier                AS [Carrier],
    o.bol_number             AS [BOL Number],
    od.item_description      AS [Item Description],
    od.lot_number            AS [Lot Number],
    o.created_by             AS [Created By],
    ISNULL(CAST(o.bol_number    AS NVARCHAR(200)), '') + ' '
     + ISNULL(CAST(od.item_number AS NVARCHAR(200)), '') + ' '
     + ISNULL(CAST(od.lot_number  AS NVARCHAR(200)), '')
                             AS [Text String],
    od.source_vessel         AS [Source Vessel],
    od.dest_vessel           AS [Destination Vessel],
    od.qty                   AS [Quantity],
    o.order_date             AS [Order Date],
    o.actual_ship_date       AS [Actual Ship Date],
    o.client_code            AS [Client ID],
    inv.invoice_number       AS [Invoice Number],
    inv.closed_date          AS [Closed Date],
    cit.contract_invoice_type_code AS [Contract Invoice Type Code],
    cb.description           AS [Description],
    -- TRIM Order: bare PW number without the warehouse prefix (e.g. "489-PW2850056" → "PW2850056")
    CASE WHEN CHARINDEX('-', o.display_order_number) > 0
         THEN SUBSTRING(o.display_order_number,
                        CHARINDEX('-', o.display_order_number) + 1,
                        LEN(o.display_order_number))
         ELSE o.display_order_number
    END                      AS [TRIM Order],
    NULL                     AS [Count Trans]
FROM Korber.t_order o
JOIN Korber.t_order_detail od
    ON o.order_id = od.order_id
LEFT JOIN Korber.t_bmm_charge c
    ON c.order_num_value = o.display_order_number
LEFT JOIN Korber.t_bmm_invoice inv
    ON c.invoice_id = inv.invoice_id
LEFT JOIN Korber.t_bmm_cont_inv_type_chargeback cb
    ON c.chargeback_id = cb.chargeback_id
LEFT JOIN Korber.t_bmm_contract_invoice_type cit
    ON cb.contract_invoice_type_id = cit.contract_invoice_type_id
WHERE o.display_order_number LIKE 'PW%'
  AND o.actual_ship_date >= ?
  AND o.actual_ship_date <  ?
"""
    df = _run_query(conn, sql, (start, end))

    # Write SSRS-format Excel: 3 preamble rows, header at row 3, data from row 4
    wb = Workbook()
    ws = wb.active
    ws.title = "HJ - Revenue Audit Packaging"

    # Row 0: empty
    ws.append([None] * len(_PKG_HEADERS))
    # Row 1: title row — "Revenue Audit Packaging" at col index 4
    title_row = [None] * len(_PKG_HEADERS)
    title_row[4] = "Revenue Audit Packaging"
    ws.append(title_row)
    # Row 2: empty
    ws.append([None] * len(_PKG_HEADERS))
    # Row 3: column headers
    ws.append(_PKG_HEADERS)

    # Data rows — map DataFrame columns to header positions
    # Build col-name → index mapping (skip None headers)
    col_idx = {h: i for i, h in enumerate(_PKG_HEADERS) if h is not None}
    for _, row in df.iterrows():
        out_row = [None] * len(_PKG_HEADERS)
        for col_name, idx in col_idx.items():
            out_row[idx] = row.get(col_name)
        ws.append(out_row)

    path = os.path.join(output_dir, "HJ - Revenue Audit Packaging.xlsx")
    wb.save(path)
    print(f"  Packaging: {len(df):,} rows", flush=True)
    return df


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Pull DC&E billing audit data from Quantix EDW (Azure SQL)"
    )
    parser.add_argument("--start", required=True, help="Start date YYYY-MM-DD (inclusive)")
    parser.add_argument("--end",   required=True, help="End date YYYY-MM-DD (exclusive)")
    parser.add_argument("--output-dir", default=".", help="Output directory (default: .)")
    args = parser.parse_args()

    start = datetime.strptime(args.start, "%Y-%m-%d").date()
    end   = datetime.strptime(args.end,   "%Y-%m-%d").date()
    output_dir = args.output_dir
    os.makedirs(output_dir, exist_ok=True)

    print("DC&E Billing Audit — EDW Pull", flush=True)
    print(f"Period : {start} to {end} (exclusive)", flush=True)
    print(f"Output : {output_dir}", flush=True)
    print("", flush=True)

    print("Connecting to EDW...", flush=True)
    try:
        conn = get_connection()
    except Exception as e:
        print(f"ERROR: Could not connect to EDW: {e}", flush=True)
        sys.exit(1)
    print("Connected.", flush=True)
    print("", flush=True)

    errors = []

    for label, fn, needs_dates in [
        ("Invoice",    lambda: pull_invoice(conn, start, end, output_dir),    True),
        ("Rate Table", lambda: pull_rate_table(conn, output_dir),             False),
        ("Lost Revenue", lambda: pull_lost_revenue(conn, start, end, output_dir), True),
        ("No Charges", lambda: pull_no_charges(conn, start, end, output_dir), True),
        ("Railcar",    lambda: pull_railcar(conn, start, end, output_dir),    True),
        ("Packaging",  lambda: pull_packaging(conn, start, end, output_dir),  True),
    ]:
        try:
            fn()
        except Exception as e:
            msg = f"ERROR — {label}: {e}"
            print(f"  {msg}", flush=True)
            errors.append(msg)

    conn.close()
    print("", flush=True)

    if errors:
        print(f"Completed with {len(errors)} error(s):", flush=True)
        for e in errors:
            print(f"  {e}", flush=True)
    else:
        print("All reports completed successfully.", flush=True)


if __name__ == "__main__":
    main()
